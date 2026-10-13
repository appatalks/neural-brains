"""Prepare an external, provenance-bound worm connectome/activity view bundle."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from numbers import Real


ROOT = Path(__file__).resolve().parents[1]
CONNECTOME_ID = "wormwiring-cook2019-corrected2020"
ACTIVITY_ID = "celegans-homogenized-calcium-activity-20250508"
CONNECTOME_FILE = "SI 5 Connectome adjacency matrices, corrected July 2020.xlsx"
ACTIVITY_FILE = "worm_data_short.parquet"
ACTIVITY_SHA256 = "17701963149851d37d40955096de181d8c5378c256fe4ede93864380f9363ddc"
MAX_OUTPUT_BYTES = 32_000_000
ACTIVITY_COLUMNS = (
    "source_dataset",
    "raw_data_file",
    "worm",
    "neuron",
    "slot",
    "is_labeled_neuron",
    "original_calcium_data",
    "original_time_in_seconds",
)


def _finite_weight(value, context: str) -> int | float:
    if value is None:
        return 0
    if isinstance(value, bool) or not isinstance(value, Real):
        raise ValueError(f"Non-numeric connectivity value in {context}: {value!r}")
    if not math.isfinite(value) or value < 0:
        raise ValueError(f"Invalid connectivity weight in {context}: {value!r}")
    return int(value) if int(value) == value else float(value)


def matrix_edges(
    row_labels: list[str],
    column_labels: list[str],
    values: list[list[object]],
    kind: str,
) -> tuple[list[dict], list[dict]]:
    if kind not in {"chemical", "electrical_gap_junction"}:
        raise ValueError("Unsupported connectome edge type")
    if len(row_labels) != len(values) or any(len(row) != len(column_labels) for row in values):
        raise ValueError("Connectivity matrix labels do not match its dimensions")
    if (
        any(not isinstance(label, str) or not label or label != label.strip() for label in row_labels + column_labels)
        or len(set(row_labels)) != len(row_labels)
        or len(set(column_labels)) != len(column_labels)
    ):
        raise ValueError("Connectivity matrix labels must be unique, non-empty strings")

    edges = []
    self_entries = []
    if kind == "electrical_gap_junction":
        if set(row_labels) != set(column_labels):
            raise ValueError("Electrical matrix must have the same row and column cell labels")
        row_index = {label: index for index, label in enumerate(row_labels)}
        column_index = {label: index for index, label in enumerate(column_labels)}
        for left in row_labels:
            for right in column_labels:
                forward = _finite_weight(
                    values[row_index[left]][column_index[right]], f"{kind} {left}->{right}"
                )
                reverse = _finite_weight(
                    values[row_index[right]][column_index[left]], f"{kind} {right}->{left}"
                )
                if forward != reverse:
                    raise ValueError(f"Electrical matrix is not symmetric at {left}, {right}")
                if left == right and forward:
                    self_entries.append({"cell": left, "weight_sections": forward})
                elif left < right and forward:
                    edges.append({
                        "source": left,
                        "target": right,
                        "kind": kind,
                        "weight_sections": forward,
                    })
        return edges, self_entries

    for row_index, source in enumerate(row_labels):
        for column_index, target in enumerate(column_labels):
            weight = _finite_weight(values[row_index][column_index], f"{kind} {source}->{target}")
            if weight:
                edges.append({
                    "source": source,
                    "target": target,
                    "kind": kind,
                    "weight_sections": weight,
                })
    return edges, self_entries


def _read_matrix_sheet(worksheet) -> tuple[list[str], list[str], list[list[object]]]:
    column_indices = [
        column for column in range(4, worksheet.max_column + 1)
        if worksheet.cell(3, column).value is not None
    ]
    row_indices = [
        row for row in range(4, worksheet.max_row + 1)
        if worksheet.cell(row, 3).value is not None
    ]
    columns = [worksheet.cell(3, column).value for column in column_indices]
    rows = [worksheet.cell(row, 3).value for row in row_indices]
    values = [
        [worksheet.cell(row, column).value for column in column_indices]
        for row in row_indices
    ]
    return rows, columns, values


def read_connectome(path: Path, sex: str) -> dict:
    if sex not in {"hermaphrodite", "male"}:
        raise ValueError("Sex must be 'hermaphrodite' or 'male'")
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=True)
    chemical_name = f"{sex} chemical"
    gap_name = f"{sex} gap jn symmetric"
    for required in ("TITLE AND LEGEND", chemical_name, gap_name):
        if required not in workbook.sheetnames:
            raise ValueError(f"Required workbook sheet is missing: {required}")

    chemical_rows, chemical_columns, chemical_values = _read_matrix_sheet(workbook[chemical_name])
    chemical_edges, _ = matrix_edges(
        chemical_rows, chemical_columns, chemical_values, "chemical"
    )
    gap_rows, gap_columns, gap_values = _read_matrix_sheet(workbook[gap_name])
    gap_edges, gap_self_entries = matrix_edges(
        gap_rows, gap_columns, gap_values, "electrical_gap_junction"
    )
    nodes = sorted(set(chemical_rows + chemical_columns + gap_rows + gap_columns))
    edges = sorted(
        chemical_edges + gap_edges,
        key=lambda edge: (edge["kind"], edge["source"], edge["target"]),
    )
    return {
        "sex": sex,
        "version": "Cook et al. 2019, corrected July 2020",
        "chemical_sheet": chemical_name,
        "electrical_sheet": gap_name,
        "nodes": nodes,
        "edges": edges,
        "edge_counts": {
            "chemical_directed": len(chemical_edges),
            "electrical_gap_undirected": len(gap_edges),
            "chemical_self_entries_retained": sum(
                edge["source"] == edge["target"] for edge in chemical_edges
            ),
            "electrical_diagonal_entries_excluded": len(gap_self_entries),
        },
        "electrical_diagonal_entries": gap_self_entries,
        "weight_semantics": (
            "Cook workbook weights are total EM serial sections of connectivity, "
            "incorporating synapse number and size; they are not conductance or "
            "measured physiological transmission strength."
        ),
        "cell_type_scope": (
            "Workbook endpoints include neurons and other cells/end organs; "
            "this adapter preserves exact cell labels but does not infer cell classes."
        ),
    }


def parse_recording_spec(specification: str) -> tuple[str, str]:
    filename, separator, worm = specification.rpartition("=")
    if not separator or not filename.endswith(".mat") or not worm.startswith("worm"):
        raise ValueError("Recording selection must be a MAT basename and worm ID, e.g. WT_NoStim.mat=worm7")
    if "/" in filename or "\\" in filename or "/" in worm or "\\" in worm:
        raise ValueError("Recording selection must not contain path components")
    return filename, worm


def validate_trace(times: list, values: list, cell_id: str) -> tuple[list[float], list[float]]:
    if not isinstance(times, (list, tuple)) or not isinstance(values, (list, tuple)):
        raise ValueError(f"Missing source time or calcium trace for {cell_id}")
    if len(times) < 2 or len(times) != len(values):
        raise ValueError(f"Source time/trace lengths are invalid for {cell_id}")
    validated_times = []
    validated_values = []
    for time, value in zip(times, values):
        if (
            isinstance(time, bool) or not isinstance(time, Real) or not math.isfinite(time)
            or isinstance(value, bool) or not isinstance(value, Real) or not math.isfinite(value)
        ):
            raise ValueError(f"Non-finite source time/calcium value for {cell_id}")
        validated_times.append(float(time))
        validated_values.append(float(value))
    if any(right < left for left, right in zip(validated_times, validated_times[1:])):
        raise ValueError(f"Source times must not move backwards for {cell_id}")
    return validated_times, validated_values


def read_activity(path: Path, connectome_nodes: set[str], selections: list[str]) -> list[dict]:
    import pyarrow.parquet as parquet

    selection_map = {}
    for selection in selections:
        filename, worm = parse_recording_spec(selection)
        key = (f"opensource_neural_data/Kato2015/{filename}", worm)
        if key in selection_map:
            raise ValueError(f"Duplicate activity recording selection: {selection}")
        selection_map[key] = {
            "recording_id": f"Kato2015/{filename}/{worm}",
            "source_dataset": "Kato2015",
            "raw_data_file": key[0],
            "worm": worm,
            "traces": [],
            "unlabeled_trace_rows": 0,
            "unmatched_labeled_cells": [],
        }
    if not selection_map:
        raise ValueError("At least one activity recording must be selected")

    file = parquet.ParquetFile(path)
    missing_columns = sorted(set(ACTIVITY_COLUMNS) - set(file.schema_arrow.names))
    if missing_columns:
        raise ValueError(f"Activity Parquet is missing required columns: {missing_columns}")
    traces_by_key = {key: set() for key in selection_map}
    times_by_key = {key: None for key in selection_map}
    for batch in file.iter_batches(batch_size=128, columns=list(ACTIVITY_COLUMNS)):
        batch_rows = batch.to_pydict()
        for index, source in enumerate(batch_rows["source_dataset"]):
            if source != "Kato2015":
                continue
            key = (batch_rows["raw_data_file"][index], batch_rows["worm"][index])
            recording = selection_map.get(key)
            if recording is None:
                continue
            if batch_rows["is_labeled_neuron"][index] is not True:
                recording["unlabeled_trace_rows"] += 1
                continue
            neuron = batch_rows["neuron"][index]
            if not isinstance(neuron, str) or not neuron:
                recording["unlabeled_trace_rows"] += 1
                continue
            if neuron not in connectome_nodes:
                recording["unmatched_labeled_cells"].append(neuron)
                continue
            if neuron in traces_by_key[key]:
                raise ValueError(f"Duplicate labelled activity trace for {key}/{neuron}")
            times, values = validate_trace(
                batch_rows["original_time_in_seconds"][index],
                batch_rows["original_calcium_data"][index],
                neuron,
            )
            if times_by_key[key] is None:
                times_by_key[key] = times
            elif times_by_key[key] != times:
                raise ValueError(f"Activity samples are not aligned across cells in {recording['recording_id']}")
            traces_by_key[key].add(neuron)
            recording["traces"].append({
                "cell": neuron,
                "slot": batch_rows["slot"][index],
                "calcium": values,
            })

    recordings = []
    for key, recording in selection_map.items():
        if not recording["traces"]:
            raise ValueError(f"No labelled traces matched the selected connectome for {key}")
        recording["traces"].sort(key=lambda trace: trace["cell"])
        recording["time_s"] = times_by_key[key]
        recording["sample_count"] = len(recording["time_s"])
        recording["unmatched_labeled_cells"] = sorted(set(recording["unmatched_labeled_cells"]))
        recording["matched_trace_count"] = len(recording["traces"])
        recording["unmatched_labeled_count"] = len(recording["unmatched_labeled_cells"])
        recording["sex_stage_metadata"] = (
            "Not encoded per recording in the processed activity table; "
            "no exact sex/stage match is asserted."
        )
        recording["measurement"] = (
            "The dataset's original_calcium_data values, not voltage or action-potential times. "
            "Samples are aligned by recorded row order; source timestamps may repeat and are not "
            "interpolated. The traces are displayed, not propagated through the connectome."
        )
        recordings.append(recording)
    recordings.sort(key=lambda recording: recording["recording_id"])
    return recordings


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _asset_path(data_dir: Path, dataset_id: str, filename: str) -> tuple[Path, dict]:
    asset_dir = data_dir / dataset_id / f"{filename}.download"
    asset = asset_dir / filename
    receipt = json.loads((asset_dir / "receipt.json").read_text(encoding="utf-8"))
    if receipt.get("dataset_id") != dataset_id or receipt.get("filename") != filename:
        raise ValueError(f"Dataset receipt does not match {asset}")
    if not asset.is_file() or asset.stat().st_size != receipt.get("size_bytes"):
        raise ValueError(f"Dataset asset does not match its recorded size: {asset}")
    if sha256_file(asset) != receipt.get("sha256"):
        raise ValueError(f"Dataset asset does not match its receipt SHA-256: {asset}")
    return asset, receipt


def prepare_bundle(data_dir: Path, sex: str, selections: list[str]) -> dict:
    data_dir = data_dir.expanduser().resolve()
    if data_dir == ROOT or ROOT in data_dir.parents:
        raise ValueError("Worm datasets must be stored outside the repository")
    connectome_path, connectome_receipt = _asset_path(
        data_dir, CONNECTOME_ID, CONNECTOME_FILE
    )
    activity_path, activity_receipt = _asset_path(
        data_dir, ACTIVITY_ID, ACTIVITY_FILE
    )
    if activity_receipt.get("upstream_checksum") != f"sha256:{ACTIVITY_SHA256}" or not activity_receipt.get("upstream_checksum_verified"):
        raise ValueError("Activity asset is not the expected provider-checksummed release")
    connectome = read_connectome(connectome_path, sex)
    recordings = read_activity(activity_path, set(connectome["nodes"]), selections)
    return {
        "schema_version": 1,
        "connectome": {
            **connectome,
            "source": "WormWiring Cook et al. corrected July 2020 workbook",
            "upstream_checksum_verified": False,
            "license_status": "User-approved local-only evaluation; upstream license remains unverified; do not redistribute.",
            "local_receipt_sha256": connectome_receipt["sha256"],
        },
        "activity": {
            "source": "qsimeon/celegans_neural_data, pinned Hugging Face release",
            "source_dataset": "Kato2015",
            "upstream_file_sha256": ACTIVITY_SHA256,
            "upstream_checksum_verified": True,
            "license": "MIT per pinned processed-dataset card; source-specific provenance retained.",
            "measurement": "calcium fluorescence",
            "recordings": recordings,
        },
        "display_limitations": [
            "The graph is a cell-label schematic, not anatomical neuron positions or reconstructed axon geometry.",
            "Calcium fluorescence playback is observational and does not represent action potentials or signal propagation.",
            "Structural contact weights are not synaptic conductances.",
            "The stylized worm motion remains kinematic and is not driven by the connectome or recordings.",
            "The processed activity table does not encode exact sex/stage for each recording.",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sex", choices=("hermaphrodite", "male"), default="hermaphrodite")
    parser.add_argument(
        "--recording",
        action="append",
        dest="recordings",
        help="Kato2015 recording as MAT basename=worm ID; repeat to include multiple recordings",
    )
    args = parser.parse_args()
    selections = args.recordings or ["WT_NoStim.mat=worm7", "WT_Stim.mat=worm0"]
    output = args.output.expanduser().resolve()
    if output == ROOT or ROOT in output.parents:
        raise ValueError("Generated worm data must stay outside the repository")
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite generated worm data: {output}")
    bundle = prepare_bundle(args.data_dir, args.sex, selections)
    payload = json.dumps(bundle, allow_nan=False, separators=(",", ":")).encode("utf-8")
    if len(payload) > MAX_OUTPUT_BYTES:
        raise ValueError(f"Generated worm bundle exceeds {MAX_OUTPUT_BYTES} bytes")
    with output.open("xb") as destination:
        destination.write(payload)
        destination.write(b"\n")
    print(f"Prepared {output} ({len(payload)} bytes)")
    print(json.dumps({
        "sex": bundle["connectome"]["sex"],
        "nodes": len(bundle["connectome"]["nodes"]),
        **bundle["connectome"]["edge_counts"],
        "recordings": [
            {
                "id": recording["recording_id"],
                "matched_traces": recording["matched_trace_count"],
                "unlabeled_rows": recording["unlabeled_trace_rows"],
                "unmatched_labeled": recording["unmatched_labeled_count"],
            }
            for recording in bundle["activity"]["recordings"]
        ],
    }, indent=2))


if __name__ == "__main__":
    main()
